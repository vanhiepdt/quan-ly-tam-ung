# -*- coding: utf-8 -*-
"""
Dựng file Google Sheets "QUẢN LÝ TẠM ỨNG - TIẾP KHÁCH - HOÀN ỨNG HÓA ĐƠN".

Nguyên tắc chống lỗi (khác hoàn toàn bản cũ):
  1. Lũy kế S/T/U dùng SUM neo đầu ($P$7:P7) thay vì chuỗi S(r-1)+...
     -> chèn/xóa dòng ở bất kỳ đâu cũng không đứt chuỗi, dòng trống vẫn đúng.
  2. Dòng Tổng cộng dùng SUM(I$7:INDEX(I:I,ROW()-1)) -> vùng tự nới khi chèn dòng.
  3. Khối KPI nằm ở SHEET RIÊNG, tham chiếu vùng tuyệt đối $7:$1000 hoặc đọc
     dòng Tổng cộng -> chèn/xóa dòng không bao giờ làm sai KPI.
  4. Toàn bộ công thức viết bằng dấu phẩy (chuẩn xlsx/US); Google Sheets tự
     chuyển sang dấu phân cách của locale khi import -> không còn #ERROR!.
"""

import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

OUT = 'QUAN LY TAM UNG - TIEP KHACH.xlsx'

S_DASH = 'Tổng quan'
S_DATA = 'Nhật ký'
S_LIST = 'Danh mục'
S_HELP = 'Hướng dẫn'

HEADER_ROW = 6
FIRST_ROW = 7
MAX_VALID = 1000          # vùng dropdown + vùng KPI quét tới dòng này

# ---------------------------------------------------------------- bảng màu
INK      = '0F172A'
INK_SOFT = '334155'
MUTED    = '64748B'
LINE     = 'CBD5E1'
LINE_SOFT= 'E2E8F0'
WHITE    = 'FFFFFF'
STRIPE   = 'F8FAFC'

GROUPS = [
    # (tiêu đề nhóm, màu nhấn, màu nền thẻ)
    ('TẠM ỨNG TỪ CƠ QUAN',        '4338CA', 'EEF2FF'),
    ('QUỸ PHÍ HÓA ĐƠN (CỦA TÔI)', '0F766E', 'F0FDFA'),
    ('QUỸ CHỊ THÚY',              'B45309', 'FFFBEB'),
    ('CHI TIẾP KHÁCH',            'BE123C', 'FFF1F2'),
    ('CHỜ XỬ LÝ',                 '475569', 'F1F5F9'),
]

# --------------------------------------------- vị trí ô giá trị của thẻ KPI
# Bố cục tab Tổng quan: mỗi nhóm chiếm 5 dòng (tiêu đề / nhãn / giá trị / ghi chú / trống),
# nhóm đầu bắt đầu ở dòng 7 nên ô giá trị của nhóm g nằm ở dòng 9 + 5*g.
# Bốn thẻ trong một nhóm bắt đầu ở các cột B, F, J, N.
CARD_COLS = ('B', 'F', 'J', 'N')


def card(group, idx):
    """Ref tuyệt đối tới ô giá trị của thẻ thứ idx trong nhóm group."""
    return f'${CARD_COLS[idx]}${9 + 5 * group}'


OK_BG, OK_FG    = 'DCFCE7', '166534'
WARN_BG, WARN_FG= 'FEF3C7', '92400E'
BAD_BG, BAD_FG  = 'FEE2E2', '991B1B'

MONEY = '#,##0" ₫";[Red]-#,##0" ₫"'
DATEF = 'dd/mm/yyyy'

# --------------------------------------------------------- cột bảng chi tiết
# (chữ cái, tiêu đề, độ rộng, kiểu)
COLS = [
    ('B', 'Ngày tháng',                             12, 'date'),
    ('C', 'Nội dung',                               46, 'text'),
    ('D', 'Ký hiệu HĐ',                             12, 'text'),
    ('E', 'Số HĐ',                                  10, 'text'),
    ('F', 'Loại HĐ',                                22, 'text'),
    ('G', 'Trạng thái HĐ đỏ',                       17, 'text'),
    ('H', 'Hình thức thanh toán',                   21, 'text'),
    ('I', 'Tổng tiền thanh toán',                   17, 'money'),
    ('J', 'Tiền rượu bia',                          15, 'money'),
    ('K', 'Hoàn tạm ứng',                           16, 'formula'),
    ('L', 'Cơ quan trả thẳng',                      17, 'formula'),
    ('M', 'Phí lấy HĐ',                             14, 'formula'),
    ('N', 'Người lấy HĐ',                           20, 'text'),
    ('O', 'Trạng thái thanh toán phí',              22, 'text'),
    ('P', 'Tạm ứng từ cơ quan',                     17, 'money'),
    ('Q', 'Giao tiền chị Thúy',                     17, 'money'),
    ('R', 'Hoàn ứng tiền mặt thừa về CQ',           20, 'money'),
    ('S', 'Dư tạm ứng lý thuyết (CQ)',              19, 'formula'),
    ('T', 'Dư tạm ứng thực tế (Hiệp)',              19, 'formula'),
    ('U', 'Dư nợ tạm ứng thực tế đang cầm (Hiệp)',  21, 'formula'),
    ('V', 'Ghi chú',                                34, 'text'),
]

DS_TRANG_THAI = ['Hợp lệ', 'Chờ HĐ', 'Không hợp lệ']
DS_HINH_THUC  = ['Tạm ứng từ cơ quan', 'Tạm ứng thêm', 'Giao tiền chị Thúy',
                 'Hoàn tạm ứng', 'Cơ quan trả thẳng', 'Nộp hoàn CQ']
DS_TT_PHI     = ['Đã thanh toán', 'Chưa thanh toán', 'Không phát sinh', 'Đã hoàn trả']
DS_LOAI_HD    = ['Hóa đơn Giá trị gia tăng', 'Hóa đơn Bán hàng', 'Phiếu chi CQ',
                 'Giấy biên nhận', 'Giấy nộp tiền', 'Khác']

TY_LE_PHI = 0.15


def serial(n):
    """Đổi serial ngày của Excel sang date."""
    return (dt.datetime(1899, 12, 30) + dt.timedelta(days=n)).date()


# Dữ liệu thật lấy từ file cũ của người dùng (bỏ dòng rác 14).
ROWS = [
    dict(ngay=serial(46269), noiDung='Nhận tạm ứng kinh phí tiếp khách từ cơ quan (đợt 1)',
         kyHieu='-', soHD='-', loaiHD='Phiếu chi CQ', trangThai='Hợp lệ',
         hinhThuc='Tạm ứng từ cơ quan', bill=0, ruou=0, nguoi='Thủ quỹ CQ',
         ttPhi='Không phát sinh', tuCQ=20000000, thuy=0, hoanCQ=0,
         ghiChu='Tiền về tay tôi, chưa giao ai'),
    dict(ngay=serial(46269), noiDung='Giao tiền tạm ứng cơ quan cho chị Thúy',
         kyHieu='-', soHD='-', loaiHD='Giấy biên nhận', trangThai='Hợp lệ',
         hinhThuc='Giao tiền chị Thúy', bill=0, ruou=0, nguoi='Chị Thúy',
         ttPhi='Không phát sinh', tuCQ=0, thuy=15000000, hoanCQ=0,
         ghiChu='Chị Thúy giữ tiền để chi tiếp khách'),
    dict(ngay=serial(46271), noiDung='Khoa sau đại học, Học viện Ngân hàng',
         kyHieu='1C26MTT', soHD='1328', loaiHD='Hóa đơn Giá trị gia tăng',
         trangThai='Hợp lệ', hinhThuc='Hoàn tạm ứng', bill=5647000, ruou=0,
         nguoi='Ngô Anh Phương', ttPhi='Chưa thanh toán', tuCQ=0, thuy=0, hoanCQ=0,
         ghiChu=''),
    dict(ngay=serial(46273),
         noiDung='Trường Đào tạo Cán bộ - Ngân hàng TMCP Đầu tư và Phát triển Việt Nam',
         kyHieu='1C26MYY', soHD='4589', loaiHD='Hóa đơn Giá trị gia tăng',
         trangThai='Hợp lệ', hinhThuc='Hoàn tạm ứng', bill=3599000, ruou=0,
         nguoi='Vương Trọng Việt Anh', ttPhi='Chưa thanh toán', tuCQ=0, thuy=0, hoanCQ=0,
         ghiChu=''),
    dict(ngay=serial(46275),
         noiDung='Viện Xã hội học và Phát triển, Học viện Chính trị Quốc gia Hồ Chí Minh',
         kyHieu='2C26MYY', soHD='2389', loaiHD='Hóa đơn Bán hàng',
         trangThai='Hợp lệ', hinhThuc='Cơ quan trả thẳng', bill=1472000, ruou=0,
         nguoi='Tự lấy trực tiếp', ttPhi='Không phát sinh', tuCQ=0, thuy=0, hoanCQ=0,
         ghiChu='Cơ quan chuyển khoản trực tiếp cho nhà cung cấp'),
    dict(ngay=serial(46280), noiDung='Cơ quan duyệt cấp bổ sung tạm ứng kinh phí tiếp khách',
         kyHieu='-', soHD='-', loaiHD='Phiếu chi CQ', trangThai='Hợp lệ',
         hinhThuc='Tạm ứng thêm', bill=0, ruou=0, nguoi='Thủ quỹ CQ',
         ttPhi='Không phát sinh', tuCQ=0, thuy=0, hoanCQ=0,
         ghiChu='Dòng mẫu - điền số tiền vào cột "Tạm ứng từ cơ quan"'),
    dict(ngay=serial(46281), noiDung='Nộp trả tiền mặt còn thừa từ quỹ phí HĐ về thủ quỹ CQ',
         kyHieu='-', soHD='-', loaiHD='Giấy nộp tiền', trangThai='Hợp lệ',
         hinhThuc='Nộp hoàn CQ', bill=0, ruou=0, nguoi='Hiệp (Tự nộp)',
         ttPhi='Đã hoàn trả', tuCQ=0, thuy=0, hoanCQ=0,
         ghiChu='Dòng mẫu - điền số tiền vào cột "Hoàn ứng tiền mặt thừa về CQ"'),
]

TOTAL_ROW = FIRST_ROW + len(ROWS)


# ------------------------------------------------------------------ tiện ích
def fill(hexv):
    return PatternFill('solid', fgColor=hexv)


def side(hexv=LINE_SOFT, style='thin'):
    return Side(style=style, color=hexv)


def box(hexv=LINE_SOFT, style='thin'):
    s = side(hexv, style)
    return Border(left=s, right=s, top=s, bottom=s)


def put(ws, ref, value, *, font=None, fillc=None, align=None, fmt=None, border=None):
    c = ws[ref]
    c.value = value
    if font:   c.font = font
    if fillc:  c.fill = fill(fillc)
    if align:  c.alignment = align
    if fmt:    c.number_format = fmt
    if border: c.border = border
    return c


def paint(ws, r1, c1, r2, c2, *, fillc=None, border=None, font=None, align=None, fmt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if fillc:  cell.fill = fill(fillc)
            if border: cell.border = border
            if font:   cell.font = font
            if align:  cell.alignment = align
            if fmt:    cell.number_format = fmt


CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left', vertical='center')
RIGHT  = Alignment(horizontal='right', vertical='center')
WRAPC  = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAPL  = Alignment(horizontal='left', vertical='center', wrap_text=True)


# ================================================================ SHEET DATA
def build_data(ws):
    ws.sheet_properties.tabColor = '1D4ED8'
    ws.sheet_view.showGridLines = False

    col = {name: letter for letter, name, _, _ in COLS}
    idx = {letter: i + 2 for i, (letter, _, _, _) in enumerate(COLS)}  # B -> 2

    ws.column_dimensions['A'].width = 2.5
    for letter, _, width, _ in COLS:
        ws.column_dimensions[letter].width = width

    LAST = COLS[-1][0]
    LASTI = idx[LAST]

    # ---- tiêu đề trang
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 34
    ws.merge_cells(f'B2:{LAST}2')
    put(ws, 'B2', 'NHẬT KÝ TẠM ỨNG, TIẾP KHÁCH & HOÀN ỨNG HÓA ĐƠN',
        font=Font(name='Calibri', size=17, bold=True, color=WHITE), align=LEFT)
    paint(ws, 2, 2, 2, LASTI, fillc=INK)
    ws['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)

    ws.row_dimensions[3].height = 22
    ws.merge_cells(f'B3:{LAST}3')
    put(ws, 'B3',
        'Chỉ nhập các cột nền trắng. Cột nền xám là công thức tự tính, đừng gõ vào. '
        'Thêm/xóa dòng bằng menu ⚙️ Quản lý tạm ứng.',
        font=Font(size=10, italic=True, color='CBD5E1'), align=LEFT)
    paint(ws, 3, 2, 3, LASTI, fillc=INK_SOFT)
    ws['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)

    ws.row_dimensions[4].height = 10
    ws.row_dimensions[5].height = 22

    # ---- dải nhãn nhóm cột trên dòng 5
    bands = [
        ('B', 'F', 'CHỨNG TỪ',           'E2E8F0', INK_SOFT),
        ('G', 'H', 'PHÂN LOẠI',          'DBEAFE', '1E40AF'),
        ('I', 'L', 'SỐ TIỀN HÓA ĐƠN',    'FFE4E6', BAD_FG),
        ('M', 'O', 'PHÍ LẤY HÓA ĐƠN',    'CCFBF1', OK_FG),
        ('P', 'R', 'DÒNG TIỀN TẠM ỨNG',  'FEF3C7', WARN_FG),
        ('S', 'U', 'SỐ DƯ LŨY KẾ',       'E0E7FF', '3730A3'),
        ('V', 'V', '',                   'F1F5F9', MUTED),
    ]
    for c1, c2, label, bg, fg in bands:
        if c1 != c2:
            ws.merge_cells(f'{c1}5:{c2}5')
        put(ws, f'{c1}5', label,
            font=Font(size=9, bold=True, color=fg), fillc=bg, align=CENTER)
        paint(ws, 5, idx[c1], 5, idx[c2], fillc=bg, border=box(WHITE))

    # ---- dòng tiêu đề cột
    ws.row_dimensions[HEADER_ROW].height = 46
    for letter, title, _, kind in COLS:
        put(ws, f'{letter}{HEADER_ROW}', title,
            font=Font(size=10, bold=True, color=WHITE),
            fillc='1E293B' if kind != 'formula' else '3730A3',
            align=WRAPC, border=box('475569'))

    # ---- dữ liệu
    for i, row in enumerate(ROWS):
        r = FIRST_ROW + i
        write_data_row(ws, r, row)

    # ---- dòng Tổng cộng
    write_total_row(ws, TOTAL_ROW, LAST, LASTI, idx)

    # ---- dropdown, định dạng có điều kiện, đóng băng
    add_validation(ws)
    add_cf(ws, LAST)

    ws.freeze_panes = f'D{FIRST_ROW}'
    ws.sheet_view.zoomScale = 100


def data_formulas(r):
    """Công thức của 5 cột tính toán trên dòng r."""
    hop_le = 'Hợp lệ'
    chi = f'SUM($I{r})-SUM($J{r})'
    return {
        'K': f'=IF(AND($G{r}="{hop_le}",$H{r}="Hoàn tạm ứng"),{chi},0)',
        'L': f'=IF(AND($G{r}="{hop_le}",$H{r}="Cơ quan trả thẳng"),{chi},0)',
        'M': f'=ROUND(SUM($K{r})*{TY_LE_PHI},0)',
        # Lũy kế bằng SUM neo đầu -> chèn/xóa dòng không đứt chuỗi
        'S': f'=SUM($P$7:$P{r})-SUM($K$7:$K{r})-SUM($R$7:$R{r})',
        'T': f'=SUM($P$7:$P{r})-SUM($Q$7:$Q{r})-SUM($M$7:$M{r})-SUM($R$7:$R{r})',
        'U': (f'=SUM($P$7:$P{r})-SUM($Q$7:$Q{r})'
              f'-SUMIF($O$7:$O{r},"Đã thanh toán",$M$7:$M{r})-SUM($R$7:$R{r})'),
    }


def write_data_row(ws, r, row):
    ws.row_dimensions[r].height = 22
    stripe = STRIPE if (r - FIRST_ROW) % 2 else WHITE
    f = data_formulas(r)

    vals = {
        'B': row['ngay'], 'C': row['noiDung'], 'D': row['kyHieu'], 'E': row['soHD'],
        'F': row['loaiHD'], 'G': row['trangThai'], 'H': row['hinhThuc'],
        'I': row['bill'], 'J': row['ruou'], 'K': f['K'], 'L': f['L'], 'M': f['M'],
        'N': row['nguoi'], 'O': row['ttPhi'], 'P': row['tuCQ'], 'Q': row['thuy'],
        'R': row['hoanCQ'], 'S': f['S'], 'T': f['T'], 'U': f['U'], 'V': row['ghiChu'],
    }

    for letter, _, _, kind in COLS:
        cell = ws[f'{letter}{r}']
        cell.value = vals[letter]
        cell.border = box(LINE_SOFT)

        if kind == 'formula':
            cell.fill = fill('EEF2FF' if letter in ('S', 'T', 'U') else 'F1F5F9')
            cell.font = Font(size=10, bold=letter in ('S', 'T', 'U'), color='1E1B4B')
            cell.number_format = MONEY
            cell.alignment = RIGHT
        elif kind == 'money':
            cell.fill = fill(stripe)
            cell.font = Font(size=10, color=INK)
            cell.number_format = MONEY
            cell.alignment = RIGHT
        elif kind == 'date':
            cell.fill = fill(stripe)
            cell.font = Font(size=10, color=INK)
            cell.number_format = DATEF
            cell.alignment = CENTER
        else:
            cell.fill = fill(stripe)
            cell.font = Font(size=10, color=INK)
            cell.alignment = WRAPL if letter in ('C', 'V') else LEFT
            if letter in ('D', 'E', 'G', 'H', 'O'):
                cell.alignment = CENTER


def write_total_row(ws, tr, LAST, LASTI, idx):
    ws.row_dimensions[tr].height = 30

    put(ws, f'B{tr}', 'Tổng cộng',
        font=Font(size=11, bold=True, color=WHITE), align=CENTER)
    ws.merge_cells(f'C{tr}:H{tr}')
    put(ws, f'C{tr}', 'TỔNG CỘNG TOÀN BỘ PHÁT SINH',
        font=Font(size=11, bold=True, color=WHITE), align=LEFT)

    # Vùng SUM tự nới khi chèn dòng ngay trên dòng này
    for letter in ['I', 'J', 'K', 'L', 'M', 'P', 'Q', 'R']:
        put(ws, f'{letter}{tr}', f'=SUM({letter}${FIRST_ROW}:INDEX({letter}:{letter},ROW()-1))',
            font=Font(size=11, bold=True, color=WHITE), align=RIGHT, fmt=MONEY)

    # Số dư cuối kỳ = ô ngay phía trên
    for letter in ['S', 'T', 'U']:
        put(ws, f'{letter}{tr}', f'=INDEX({letter}:{letter},ROW()-1)',
            font=Font(size=11, bold=True, color='FDE68A'), align=RIGHT, fmt=MONEY)

    paint(ws, tr, 2, tr, LASTI, border=box('0F172A'))
    for c in range(2, LASTI + 1):
        ws.cell(row=tr, column=c).fill = fill(INK)
    ws[f'B{tr}'].font = Font(size=11, bold=True, color=WHITE)
    ws[f'C{tr}'].font = Font(size=11, bold=True, color=WHITE)

    # chân trang nhắc việc
    fr = tr + 2
    ws.merge_cells(f'B{fr}:{LAST}{fr}')
    put(ws, f'B{fr}',
        'Mẹo: cần thêm dòng thì bấm menu ⚙️ Quản lý tạm ứng → ➕ Thêm 1 dòng mới. '
        'Nếu thấy số lạ, bấm 🔧 Sửa toàn bộ công thức.',
        font=Font(size=9, italic=True, color=MUTED), align=LEFT)


def add_validation(ws):
    specs = [
        ('G', DS_TRANG_THAI, 'Trạng thái hóa đơn đỏ',
         'Chọn: Hợp lệ / Chờ HĐ / Không hợp lệ'),
        ('H', DS_HINH_THUC, 'Hình thức thanh toán',
         'Chọn đúng 1 hình thức trong danh sách'),
        ('O', DS_TT_PHI, 'Trạng thái thanh toán phí',
         'Chọn: Đã thanh toán / Chưa thanh toán / Không phát sinh / Đã hoàn trả'),
        ('F', DS_LOAI_HD, 'Loại hóa đơn / chứng từ',
         'Chọn loại chứng từ'),
    ]
    for letter, items, title, msg in specs:
        dv = DataValidation(
            type='list',
            formula1='"' + ','.join(items) + '"',
            allow_blank=True,
            showDropDown=False,   # False = HIỆN mũi tên dropdown (quirk của xlsx)
        )
        dv.error = 'Giá trị không nằm trong danh sách cho phép.'
        dv.errorTitle = 'Sai giá trị'
        dv.prompt = msg
        dv.promptTitle = title
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f'{letter}{FIRST_ROW}:{letter}{MAX_VALID}')


def add_cf(ws, LAST):
    end = MAX_VALID

    def rng(letter):
        return f'{letter}{FIRST_ROW}:{letter}{end}'

    # Trạng thái HĐ đỏ
    ws.conditional_formatting.add(rng('G'), FormulaRule(
        formula=[f'$G{FIRST_ROW}="Hợp lệ"'], fill=fill(OK_BG),
        font=Font(color=OK_FG, bold=True), stopIfTrue=False))
    ws.conditional_formatting.add(rng('G'), FormulaRule(
        formula=[f'$G{FIRST_ROW}="Chờ HĐ"'], fill=fill(WARN_BG),
        font=Font(color=WARN_FG, bold=True), stopIfTrue=False))
    ws.conditional_formatting.add(rng('G'), FormulaRule(
        formula=[f'$G{FIRST_ROW}="Không hợp lệ"'], fill=fill(BAD_BG),
        font=Font(color=BAD_FG, bold=True), stopIfTrue=False))

    # Trạng thái thanh toán phí
    ws.conditional_formatting.add(rng('O'), FormulaRule(
        formula=[f'$O{FIRST_ROW}="Đã thanh toán"'], fill=fill(OK_BG),
        font=Font(color=OK_FG, bold=True), stopIfTrue=False))
    ws.conditional_formatting.add(rng('O'), FormulaRule(
        formula=[f'$O{FIRST_ROW}="Chưa thanh toán"'], fill=fill(BAD_BG),
        font=Font(color=BAD_FG, bold=True), stopIfTrue=False))
    ws.conditional_formatting.add(rng('O'), FormulaRule(
        formula=[f'$O{FIRST_ROW}="Đã hoàn trả"'], fill=fill('DBEAFE'),
        font=Font(color='1E40AF', bold=True), stopIfTrue=False))

    # Hình thức thanh toán: tô nhạt theo nhóm dòng tiền
    ws.conditional_formatting.add(rng('H'), FormulaRule(
        formula=[f'OR($H{FIRST_ROW}="Tạm ứng từ cơ quan",$H{FIRST_ROW}="Tạm ứng thêm")'],
        fill=fill('FEF3C7'), font=Font(color=WARN_FG, bold=True), stopIfTrue=False))
    ws.conditional_formatting.add(rng('H'), FormulaRule(
        formula=[f'$H{FIRST_ROW}="Nộp hoàn CQ"'],
        fill=fill('E0E7FF'), font=Font(color='3730A3', bold=True), stopIfTrue=False))

    # Số dư âm -> đỏ
    for letter in ('S', 'T', 'U'):
        ws.conditional_formatting.add(rng(letter), CellIsRule(
            operator='lessThan', formula=['0'],
            fill=fill(BAD_BG), font=Font(color=BAD_FG, bold=True)))

    # Dòng chưa nhập nội dung -> làm mờ để dễ thấy chỗ cần điền
    ws.conditional_formatting.add(f'C{FIRST_ROW}:C{end}', FormulaRule(
        formula=[f'AND($C{FIRST_ROW}="",$B{FIRST_ROW}<>"")'],
        fill=fill('FEF2F2'), stopIfTrue=False))


# =========================================================== SHEET DASHBOARD
def kpi_specs():
    """(nhóm, nhãn, công thức, ghi chú ngắn)"""
    D = f"'{S_DATA}'"
    TR = TOTAL_ROW

    def col(letter):                      # vùng tuyệt đối, không xê dịch khi chèn dòng
        return f'{D}!${letter}${FIRST_ROW}:${letter}${MAX_VALID}'

    def tot(letter):                      # đọc dòng Tổng cộng (tự đi theo khi dòng dịch)
        return f'{D}!{letter}{TR}'

    thuy_goc = (f'IFERROR(INDEX({col("Q")},MATCH("Giao tiền chị Thúy",{col("H")},0)),0)')

    return [
        (0, 'Tổng tạm ứng từ cơ quan', f'={tot("P")}',
         'Tổng tiền cơ quan đã ứng ra'),
        (0, 'Cơ quan tạm ứng thêm',
         f'=SUMIF({col("H")},"Tạm ứng thêm",{col("P")})',
         'Phần ứng bổ sung sau đợt đầu'),
        (0, 'Tổng tồn quỹ cơ quan', f'={card(1, 3)}+{card(2, 3)}',
         'Tiền cơ quan còn nằm ở 2 quỹ'),
        (0, 'Dư nợ CQ (lý thuyết)', f'={tot("S")}',
         'Số phải hoàn theo sổ cơ quan'),

        (1, 'Quỹ tôi nhận (phí HĐ)', f'={tot("P")}-{tot("Q")}',
         'Tạm ứng giữ lại, không giao chị Thúy'),
        (1, 'Quỹ tôi: đã chi trả phí',
         f'=SUMIF({col("O")},"Đã thanh toán",{col("M")})',
         'Chỉ tính dòng đã thanh toán'),
        (1, 'Quỹ tôi: nộp trả tiền thừa', f'={tot("R")}',
         'Đã nộp lại thủ quỹ cơ quan'),
        (1, 'Quỹ tôi: dư còn lại',
         f'={card(1, 0)}-{card(1, 1)}-{card(1, 2)}',
         'Tiền mặt đang cầm của quỹ phí'),

        (2, 'Giao chị Thúy (gốc)', f'={thuy_goc}',
         'Lần giao đầu tiên'),
        (2, 'Giao thêm chị Thúy', f'={tot("Q")}-{card(2, 0)}',
         'Các lần giao sau'),
        (2, 'Chị Thúy: đã hoàn ứng', f'={tot("K")}',
         'Bill hợp lệ đã hoàn tạm ứng'),
        (2, 'Chị Thúy: dư còn lại',
         f'={card(2, 0)}+{card(2, 1)}-{card(2, 2)}',
         'Chị Thúy còn giữ bao nhiêu'),

        (3, 'Tổng chi tiếp khách (bill)', f'={tot("I")}',
         'Tổng giá trị hóa đơn'),
        (3, 'Tiền rượu bia (không hoàn)', f'={tot("J")}',
         'Phần bị loại khỏi hoàn ứng'),
        (3, 'Cơ quan trả thẳng (đã trả)', f'={tot("L")}',
         'Cơ quan trả trực tiếp NCC'),
        (3, 'Tổng chi thực tế 2 quỹ', f'={card(1, 1)}+{card(2, 2)}',
         'Phí đã trả + hoàn ứng cho Thúy'),

        (4, 'Tiền chờ HĐ (tạm ứng)',
         f'=SUMIFS({col("I")},{col("G")},"Chờ HĐ",{col("H")},"Hoàn tạm ứng")'
         f'-SUMIFS({col("J")},{col("G")},"Chờ HĐ",{col("H")},"Hoàn tạm ứng")',
         'Bill chưa có hóa đơn đỏ'),
        (4, 'Chờ CQ trả thẳng',
         f'=SUMIFS({col("I")},{col("G")},"Chờ HĐ",{col("H")},"Cơ quan trả thẳng")'
         f'-SUMIFS({col("J")},{col("G")},"Chờ HĐ",{col("H")},"Cơ quan trả thẳng")',
         'Chờ cơ quan chuyển cho NCC'),
        (4, 'Dư thực tế (Hiệp)', f'={tot("T")}',
         'Đã trừ toàn bộ phí lấy HĐ'),
        (4, 'Dư thực tế đang cầm (Hiệp)', f'={tot("U")}',
         'Chỉ trừ phí đã thanh toán'),
    ]


def build_dash(ws):
    ws.sheet_properties.tabColor = '0F172A'
    ws.sheet_view.showGridLines = False

    # A spacer | 4 thẻ, mỗi thẻ 3 cột, giữa các thẻ 1 cột hẹp
    ws.column_dimensions['A'].width = 3
    card_cols = []                       # cột bắt đầu của từng thẻ
    c = 2
    for i in range(4):
        card_cols.append(c)
        for k in range(3):
            ws.column_dimensions[get_column_letter(c + k)].width = 13.5
        c += 3
        if i < 3:
            ws.column_dimensions[get_column_letter(c)].width = 2.5
            c += 1
    LASTI = c - 2
    LAST = get_column_letter(LASTI)

    # ---- banner
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 40
    ws.merge_cells(f'B2:{LAST}2')
    put(ws, 'B2', 'BẢNG THEO DÕI TẠM ỨNG, TIẾP KHÁCH & HOÀN ỨNG HÓA ĐƠN',
        font=Font(size=19, bold=True, color=WHITE))
    ws['B2'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    paint(ws, 2, 2, 2, LASTI, fillc=INK)

    ws.row_dimensions[3].height = 24
    ws.merge_cells(f'B3:{LAST}3')
    put(ws, 'B3', 'Số liệu tự cập nhật từ tab "Nhật ký". Không nhập tay vào trang này.',
        font=Font(size=10, color='94A3B8'))
    ws['B3'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    paint(ws, 3, 2, 3, LASTI, fillc=INK_SOFT)

    # ---- banner đối soát
    ws.row_dimensions[4].height = 10
    ws.row_dimensions[5].height = 40
    ws.merge_cells(f'B5:{LAST}5')
    put(ws, 'B5',
        f'=IF(ROUND({card(0, 2)}+{card(3, 3)}+{card(1, 2)},0)'
        f'=ROUND({card(0, 0)},0),'
        '"✅  ĐÚNG CÂN ĐỐI  —  Tồn quỹ + đã chi + đã nộp trả = tổng tạm ứng",'
        '"⚠️  LỆCH QUỸ  —  kiểm tra lại các cột số tiền ở tab Nhật ký")',
        font=Font(size=13, bold=True), align=CENTER)
    paint(ws, 5, 2, 5, LASTI, border=box(LINE))
    ws.conditional_formatting.add(f'B5:{LAST}5', FormulaRule(
        formula=['LEFT($B$5,1)="✅"'], fill=fill(OK_BG), font=Font(color=OK_FG, bold=True)))
    ws.conditional_formatting.add(f'B5:{LAST}5', FormulaRule(
        formula=['LEFT($B$5,1)="⚠"'], fill=fill(WARN_BG), font=Font(color=WARN_FG, bold=True)))

    # ---- các nhóm thẻ
    specs = kpi_specs()
    row = 7
    positions = {}                       # (nhóm, chỉ số trong nhóm) -> ô giá trị

    for gi, (gname, accent, bg) in enumerate(GROUPS):
        # tiêu đề nhóm
        ws.row_dimensions[row].height = 26
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=LASTI)
        cell = ws.cell(row=row, column=2)
        cell.value = f'  {gname}'
        cell.font = Font(size=11, bold=True, color=WHITE)
        cell.alignment = LEFT
        paint(ws, row, 2, row, LASTI, fillc=accent)

        items = [s for s in specs if s[0] == gi]
        r_label, r_value, r_note = row + 1, row + 2, row + 3
        ws.row_dimensions[r_label].height = 20
        ws.row_dimensions[r_value].height = 30
        ws.row_dimensions[r_note].height = 18

        for ci, (_, label, formula, note) in enumerate(items):
            c0 = card_cols[ci]
            c1 = c0 + 2
            L0 = get_column_letter(c0)

            ws.merge_cells(start_row=r_label, start_column=c0, end_row=r_label, end_column=c1)
            put(ws, f'{L0}{r_label}', label,
                font=Font(size=9, bold=True, color=accent), align=CENTER)

            ws.merge_cells(start_row=r_value, start_column=c0, end_row=r_value, end_column=c1)
            put(ws, f'{L0}{r_value}', formula,
                font=Font(size=14, bold=True, color=INK), align=CENTER, fmt=MONEY)

            ws.merge_cells(start_row=r_note, start_column=c0, end_row=r_note, end_column=c1)
            put(ws, f'{L0}{r_note}', note,
                font=Font(size=8, italic=True, color=MUTED), align=CENTER)

            paint(ws, r_label, c0, r_note, c1, fillc=bg, border=box(LINE))
            positions[(gi, ci)] = f'{L0}{r_value}'

        row = r_note + 2

    # chú thích cuối
    ws.row_dimensions[row].height = 20
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=LASTI)
    put(ws, ws.cell(row=row, column=2).coordinate,
        'Quỹ tôi = phần tạm ứng giữ lại để trả phí lấy hóa đơn. '
        'Quỹ chị Thúy = phần đã giao để chi tiếp khách.',
        font=Font(size=9, italic=True, color=MUTED), align=LEFT)

    return positions


# =============================================================== SHEET LISTS
def build_lists(ws):
    ws.sheet_properties.tabColor = '94A3B8'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    for letter, width in zip('BCDE', (26, 26, 26, 30)):
        ws.column_dimensions[letter].width = width

    ws.row_dimensions[2].height = 30
    ws.merge_cells('B2:E2')
    put(ws, 'B2', 'DANH MỤC LỰA CHỌN & THAM SỐ',
        font=Font(size=13, bold=True, color=WHITE), align=LEFT)
    paint(ws, 2, 2, 2, 5, fillc=INK)

    heads = ['Trạng thái HĐ đỏ', 'Hình thức thanh toán',
             'Trạng thái thanh toán phí', 'Loại HĐ / chứng từ']
    lists = [DS_TRANG_THAI, DS_HINH_THUC, DS_TT_PHI, DS_LOAI_HD]

    for i, h in enumerate(heads):
        put(ws, f'{get_column_letter(2 + i)}4', h,
            font=Font(size=10, bold=True, color=WHITE), fillc='1E293B',
            align=WRAPC, border=box('475569'))
    ws.row_dimensions[4].height = 32

    for i, items in enumerate(lists):
        for j, it in enumerate(items):
            put(ws, f'{get_column_letter(2 + i)}{5 + j}', it,
                font=Font(size=10), fillc=WHITE if j % 2 else STRIPE,
                align=LEFT, border=box(LINE_SOFT))

    put(ws, 'B14', 'Tỷ lệ phí lấy hóa đơn',
        font=Font(size=10, bold=True, color=WHITE), fillc='1E293B', align=LEFT,
        border=box('475569'))
    put(ws, 'C14', TY_LE_PHI, font=Font(size=12, bold=True, color=OK_FG),
        fillc=OK_BG, align=CENTER, fmt='0.0%', border=box(LINE))
    ws.merge_cells('D14:E14')
    put(ws, 'D14', 'Đổi ở đây rồi bấm 🔧 Sửa toàn bộ công thức',
        font=Font(size=9, italic=True, color=MUTED), align=LEFT)

    ws.merge_cells('B16:E16')
    put(ws, 'B16',
        'Sửa danh mục ở đây KHÔNG tự đổi dropdown trong tab Nhật ký. '
        'Muốn áp dụng, sửa mảng tương ứng trong Apps Script rồi bấm 🔧 Sửa toàn bộ công thức.',
        font=Font(size=9, italic=True, color=WARN_FG), align=WRAPL)
    ws.row_dimensions[16].height = 30


# ================================================================ SHEET HELP
HELP = [
    ('h', 'CÁCH DÙNG BẢNG NÀY', ''),
    ('t', '1. Nhập liệu ở tab "Nhật ký"',
     'Chỉ gõ vào các cột nền trắng: Ngày tháng, Nội dung, Ký hiệu HĐ, Số HĐ, Loại HĐ, '
     'Trạng thái HĐ đỏ, Hình thức thanh toán, Tổng tiền, Tiền rượu bia, Người lấy HĐ, '
     'Trạng thái thanh toán phí, Tạm ứng từ cơ quan, Giao tiền chị Thúy, '
     'Hoàn ứng tiền mặt thừa về CQ, Ghi chú.'),
    ('t', '2. Các cột KHÔNG gõ vào',
     'K (Hoàn tạm ứng), L (Cơ quan trả thẳng), M (Phí lấy HĐ), S, T, U (số dư lũy kế) '
     'và cả dòng Tổng cộng đều là công thức. Nền của chúng có màu xám hoặc tím nhạt.'),
    ('t', '3. Thêm / xóa dòng',
     'Dùng menu ⚙️ Quản lý tạm ứng → ➕ Thêm 1 dòng mới hoặc 🗑️ Xóa dòng đang chọn. '
     'Không chèn dòng bằng chuột phải: dòng chèn tay sẽ không có công thức.'),
    ('t', '4. Khi thấy số lạ hoặc ô báo lỗi',
     'Bấm ⚙️ Quản lý tạm ứng → 🔧 Sửa toàn bộ công thức. Hàm này ghi lại toàn bộ công thức '
     'theo đúng vị trí dòng hiện tại và báo lại nếu còn ô lỗi.'),
    ('t', '5. Ý nghĩa 3 cột số dư',
     'S = Dư nợ theo sổ cơ quan (tạm ứng trừ phần đã hoàn ứng và đã nộp trả). '
     'T = Tiền thực tế còn lại của tôi nếu đã trả hết phí lấy hóa đơn. '
     'U = Tiền mặt thực đang cầm, chỉ trừ phí của dòng đã đánh dấu "Đã thanh toán".'),
    ('t', '6. Ba loại "hình thức thanh toán" hay nhầm',
     '"Hoàn tạm ứng" = bill do quỹ chị Thúy chi, cơ quan hoàn lại. '
     '"Cơ quan trả thẳng" = cơ quan chuyển khoản trực tiếp cho nhà cung cấp, tiền không qua tay. '
     '"Nộp hoàn CQ" = nộp tiền mặt còn thừa về thủ quỹ, điền số vào cột R.'),
    ('t', '7. Vì sao bản này không còn lệch dòng',
     'Số dư lũy kế dùng SUM neo từ dòng đầu ($P$7:$P9) thay vì cộng dồn từ dòng trên. '
     'Chèn hay xóa dòng ở giữa bảng thì vùng tự nới, chuỗi không bao giờ đứt. '
     'Khối chỉ số nằm ở tab riêng nên cũng không bị ảnh hưởng.'),
]


def build_help(ws):
    ws.sheet_properties.tabColor = '0D9488'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 88

    r = 2
    for kind, a, b in HELP:
        if kind == 'h':
            ws.row_dimensions[r].height = 36
            ws.merge_cells(f'B{r}:C{r}')
            put(ws, f'B{r}', a, font=Font(size=16, bold=True, color=WHITE), align=LEFT)
            paint(ws, r, 2, r, 3, fillc=INK)
            r += 2
        else:
            ws.row_dimensions[r].height = 46
            put(ws, f'B{r}', a, font=Font(size=11, bold=True, color='1E293B'),
                fillc='E2E8F0', align=WRAPL, border=box(LINE))
            put(ws, f'C{r}', b, font=Font(size=10, color=INK_SOFT),
                fillc=WHITE, align=WRAPL, border=box(LINE))
            r += 1


# ===================================================================== main
def main():
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = S_DASH
    ws_data = wb.create_sheet(S_DATA)
    ws_list = wb.create_sheet(S_LIST)
    ws_help = wb.create_sheet(S_HELP)

    build_data(ws_data)
    build_dash(ws_dash)
    build_lists(ws_list)
    build_help(ws_help)

    wb.active = 0
    wb.save(OUT)
    print('saved', OUT, '| total row =', TOTAL_ROW)


if __name__ == '__main__':
    main()
